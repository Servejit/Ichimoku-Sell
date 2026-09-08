```python
import streamlit as st
import pandas as pd
import numpy as np
import yfinance as yf
import requests
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule


# ============================================================
# STREAMLIT PAGE
# ============================================================

st.set_page_config(
    page_title="6thSense — NIFTY 200 Sell Scanner",
    page_icon="📉",
    layout="wide"
)

st.title("Welcome to 6thSense Trading")
st.subheader("📉 NIFTY 200 — ICHIMOKU + VOLUME + MACD BEARISH / SELL SCANNER")


# ============================================================
# SETTINGS
# ============================================================

MAX_WORKERS = 3

PERIOD_5M = "5d"
PERIOD_15M = "5d"
PERIOD_1H = "30d"

MIN_BEARISH_SCORE = 45
MAX_MACD_SIGNAL_AGE = 3

TENKAN_PERIOD = 9
KIJUN_PERIOD = 26
SENKOU_PERIOD = 52

VOLUME_PERIOD = 20

MACD_FAST = 12
MACD_SLOW = 26
MACD_SIGNAL = 9

NIFTY_URL = (
    "https://www.niftyindices.com/"
    "IndexConstituent/ind_nifty200list.csv"
)


# ============================================================
# GET NIFTY 200
# ============================================================

@st.cache_data(ttl=3600)
def get_nifty200():

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0 Safari/537.36"
        ),
        "Accept": "text/csv,*/*",
        "Referer": "https://www.niftyindices.com/"
    }

    response = requests.get(
        NIFTY_URL,
        headers=headers,
        timeout=30
    )

    response.raise_for_status()

    nifty_df = pd.read_csv(
        BytesIO(response.content)
    )

    if "Symbol" not in nifty_df.columns:
        raise ValueError(
            "NIFTY 200 CSV does not contain Symbol column."
        )

    symbols = (
        nifty_df["Symbol"]
        .astype(str)
        .str.strip()
        .str.upper()
        .dropna()
        .drop_duplicates()
        .tolist()
    )

    return symbols


# ============================================================
# DOWNLOAD DATA
# ============================================================

def download_data(interval, period, tickers):

    data = yf.download(
        tickers=tickers,
        period=period,
        interval=interval,
        auto_adjust=False,
        group_by="ticker",
        threads=False,
        progress=False
    )

    return data


# ============================================================
# EXTRACT INDIVIDUAL STOCK DATA
# ============================================================

def get_stock_data(df, ticker):

    if df is None or df.empty:
        return pd.DataFrame()

    try:

        if isinstance(df.columns, pd.MultiIndex):

            if ticker in df.columns.get_level_values(0):

                out = df[ticker].copy()

            elif ticker in df.columns.get_level_values(1):

                out = df.xs(
                    ticker,
                    axis=1,
                    level=1
                ).copy()

            else:
                return pd.DataFrame()

        else:
            out = df.copy()

    except Exception:
        return pd.DataFrame()

    rename_map = {}

    for col in out.columns:

        c = str(col).strip().lower()

        if c == "open":
            rename_map[col] = "Open"

        elif c == "high":
            rename_map[col] = "High"

        elif c == "low":
            rename_map[col] = "Low"

        elif c == "close":
            rename_map[col] = "Close"

        elif c == "volume":
            rename_map[col] = "Volume"

    out = out.rename(
        columns=rename_map
    )

    required = [
        "Open",
        "High",
        "Low",
        "Close",
        "Volume"
    ]

    for col in required:

        if col not in out.columns:
            out[col] = np.nan

    out = out[required].copy()

    for col in required:

        out[col] = pd.to_numeric(
            out[col],
            errors="coerce"
        )

    out = out.dropna(
        subset=[
            "High",
            "Low",
            "Close"
        ]
    )

    out = out.sort_index()

    out = out[
        ~out.index.duplicated(
            keep="last"
        )
    ]

    # ========================================================
    # TIMEZONE FIX
    # ========================================================

    if isinstance(
        out.index,
        pd.DatetimeIndex
    ):

        if out.index.tz is not None:

            out.index = (
                out.index
                .tz_convert("Asia/Kolkata")
                .tz_localize(None)
            )

    return out


# ============================================================
# REMOVE INCOMPLETE LAST CANDLE
# ============================================================

def remove_incomplete_last_candle(
    df,
    interval_minutes
):

    if df.empty:
        return df

    try:

        now = (
            pd.Timestamp.now(
                tz="Asia/Kolkata"
            )
            .tz_localize(None)
        )

        last_time = df.index[-1]

        expected_end = (
            last_time
            + pd.Timedelta(
                minutes=interval_minutes
            )
        )

        if expected_end > now:

            df = df.iloc[:-1].copy()

    except Exception:
        pass

    return df


# ============================================================
# ICHIMOKU + VOLUME + MACD
# ============================================================

def add_indicators(df):

    df = df.copy()

    high = df["High"]
    low = df["Low"]

    # --------------------------------------------------------
    # ICHIMOKU
    # --------------------------------------------------------

    df["Tenkan"] = (
        high.rolling(
            TENKAN_PERIOD
        ).max()
        +
        low.rolling(
            TENKAN_PERIOD
        ).min()
    ) / 2

    df["Kijun"] = (
        high.rolling(
            KIJUN_PERIOD
        ).max()
        +
        low.rolling(
            KIJUN_PERIOD
        ).min()
    ) / 2

    df["Span_A"] = (
        df["Tenkan"]
        +
        df["Kijun"]
    ) / 2

    df["Span_B"] = (
        high.rolling(
            SENKOU_PERIOD
        ).max()
        +
        low.rolling(
            SENKOU_PERIOD
        ).min()
    ) / 2

    df["Cloud_Top"] = df[
        ["Span_A", "Span_B"]
    ].max(axis=1)

    df["Cloud_Bottom"] = df[
        ["Span_A", "Span_B"]
    ].min(axis=1)

    # --------------------------------------------------------
    # VOLUME
    # --------------------------------------------------------

    df["Volume_Avg20"] = (
        df["Volume"]
        .rolling(VOLUME_PERIOD)
        .mean()
    )

    df["Volume_Ratio"] = np.where(
        df["Volume_Avg20"] > 0,
        df["Volume"] /
        df["Volume_Avg20"],
        np.nan
    )

    df["Volume_Avg5"] = (
        df["Volume"]
        .rolling(5)
        .mean()
    )

    # --------------------------------------------------------
    # MACD
    # --------------------------------------------------------

    ema_fast = (
        df["Close"]
        .ewm(
            span=MACD_FAST,
            adjust=False
        )
        .mean()
    )

    ema_slow = (
        df["Close"]
        .ewm(
            span=MACD_SLOW,
            adjust=False
        )
        .mean()
    )

    df["MACD"] = (
        ema_fast -
        ema_slow
    )

    df["MACD_Signal"] = (
        df["MACD"]
        .ewm(
            span=MACD_SIGNAL,
            adjust=False
        )
        .mean()
    )

    df["MACD_Histogram"] = (
        df["MACD"]
        -
        df["MACD_Signal"]
    )

    # --------------------------------------------------------
    # MACD CROSS
    # --------------------------------------------------------

    df["MACD_Buy_Cross"] = (
        (
            df["MACD"].shift(1)
            <=
            df["MACD_Signal"].shift(1)
        )
        &
        (
            df["MACD"]
            >
            df["MACD_Signal"]
        )
    )

    df["MACD_Sell_Cross"] = (
        (
            df["MACD"].shift(1)
            >=
            df["MACD_Signal"].shift(1)
        )
        &
        (
            df["MACD"]
            <
            df["MACD_Signal"]
        )
    )

    return df


# ============================================================
# VOLUME SCORE
# ============================================================

def volume_score(value):

    if pd.isna(value):
        return 0

    if value >= 2:
        return 10

    elif value >= 1.5:
        return 8

    elif value >= 1.2:
        return 6

    elif value >= 1:
        return 3

    return 0


def volume_score_1h(value):

    if pd.isna(value):
        return 0

    if value >= 2:
        return 15

    elif value >= 1.5:
        return 12

    elif value >= 1.2:
        return 9

    elif value >= 1:
        return 5

    return 0


# ============================================================
# FORMAT TIME
# ============================================================

def format_time(value):

    if value is None:
        return ""

    try:

        return pd.Timestamp(
            value
        ).strftime(
            "%Y-%m-%d %H:%M"
        )

    except Exception:

        return str(value)


# ============================================================
# 5 MINUTE BEARISH ANALYSIS
# ============================================================

def analyze_5m(df):

    if df.empty or len(df) < 80:
        return None

    df = add_indicators(df)

    df = df.dropna(
        subset=[
            "Tenkan",
            "Kijun",
            "Span_A",
            "Span_B",
            "MACD",
            "MACD_Signal"
        ]
    )

    if len(df) < 2:
        return None

    current = df.iloc[-1]
    previous = df.iloc[-2]

    score = 0

    # --------------------------------------------------------
    # BEARISH ICHIMOKU
    # --------------------------------------------------------

    if current["Close"] < current["Tenkan"]:
        score += 10

    if current["Close"] < current["Kijun"]:
        score += 10

    if current["Tenkan"] < current["Kijun"]:
        score += 15

    if current["Close"] < current["Cloud_Bottom"]:
        score += 10

    if current["Span_A"] < current["Span_B"]:
        score += 5

    if current["Tenkan"] < previous["Tenkan"]:
        score += 5

    # --------------------------------------------------------
    # VOLUME
    # --------------------------------------------------------

    score += volume_score(
        current["Volume_Ratio"]
    )

    # --------------------------------------------------------
    # LAST BEARISH ICHIMOKU CROSS
    # --------------------------------------------------------

    cross = (
        (df["Tenkan"].shift(1)
         >=
         df["Kijun"].shift(1))
        &
        (df["Tenkan"]
         <
         df["Kijun"])
    )

    cross_positions = np.where(
        cross.values
    )[0]

    cross_age = None
    cross_time = None

    if len(cross_positions) > 0:

        last_cross = cross_positions[-1]

        cross_age = (
            len(df)
            - 1
            - last_cross
        )

        cross_time = df.index[
            last_cross
        ]

    # --------------------------------------------------------
    # MACD
    # --------------------------------------------------------

    sell_cross_positions = np.where(
        df["MACD_Sell_Cross"].values
    )[0]

    buy_cross_positions = np.where(
        df["MACD_Buy_Cross"].values
    )[0]

    macd_signal = "NONE"
    macd_cross_time = None
    macd_cross_age = None
    macd_just_signal = 0

    latest_sell = (
        sell_cross_positions[-1]
        if len(sell_cross_positions) > 0
        else -1
    )

    latest_buy = (
        buy_cross_positions[-1]
        if len(buy_cross_positions) > 0
        else -1
    )

    if latest_sell > latest_buy and latest_sell >= 0:

        age = (
            len(df)
            - 1
            - latest_sell
        )

        macd_cross_time = df.index[
            latest_sell
        ]

        macd_cross_age = age

        if age <= MAX_MACD_SIGNAL_AGE:

            macd_signal = "FRESH SELL"

            if age == 0:
                score += 12
                macd_just_signal = 12

            elif age == 1:
                score += 10
                macd_just_signal = 10

            elif age == 2:
                score += 8
                macd_just_signal = 8

            elif age == 3:
                score += 5
                macd_just_signal = 5

    elif latest_buy > latest_sell and latest_buy >= 0:

        age = (
            len(df)
            - 1
            - latest_buy
        )

        macd_cross_time = df.index[
            latest_buy
        ]

        macd_cross_age = age

        if age <= MAX_MACD_SIGNAL_AGE:
            macd_signal = "FRESH BUY"

    return {

        "score": score,

        "close": current["Close"],

        "tenkan": current["Tenkan"],

        "kijun": current["Kijun"],

        "span_a": current["Span_A"],

        "span_b": current["Span_B"],

        "cloud_bottom": current["Cloud_Bottom"],

        "volume": current["Volume"],

        "volume_avg20": current["Volume_Avg20"],

        "volume_ratio": current["Volume_Ratio"],

        "cross_age": cross_age,

        "cross_time": cross_time,

        "macd": current["MACD"],

        "macd_signal_value":
            current["MACD_Signal"],

        "macd_histogram":
            current["MACD_Histogram"],

        "macd_signal":
            macd_signal,

        "macd_just_signal":
            macd_just_signal,

        "macd_cross_time":
            macd_cross_time,

        "macd_cross_age":
            macd_cross_age
    }


# ============================================================
# 15 MINUTE BEARISH ANALYSIS
# ============================================================

def analyze_15m(df):

    if df.empty or len(df) < 80:
        return None

    df = add_indicators(df)

    df = df.dropna(
        subset=[
            "Tenkan",
            "Kijun",
            "Span_A",
            "Span_B"
        ]
    )

    if len(df) < 2:
        return None

    current = df.iloc[-1]
    previous = df.iloc[-2]

    score = 0

    if current["Close"] < current["Tenkan"]:
        score += 10

    if current["Close"] < current["Kijun"]:
        score += 10

    if current["Tenkan"] < current["Kijun"]:
        score += 15

    if current["Close"] < current["Cloud_Bottom"]:
        score += 15

    if current["Span_A"] < current["Span_B"]:
        score += 10

    if current["Tenkan"] < previous["Tenkan"]:
        score += 5

    if current["Kijun"] < previous["Kijun"]:
        score += 5

    score += volume_score(
        current["Volume_Ratio"]
    )

    return {

        "score": score,

        "close": current["Close"],

        "tenkan": current["Tenkan"],

        "kijun": current["Kijun"],

        "span_a": current["Span_A"],

        "span_b": current["Span_B"],

        "cloud_bottom":
            current["Cloud_Bottom"],

        "volume_ratio":
            current["Volume_Ratio"]
    }


# ============================================================
# 1 HOUR BEARISH ANALYSIS
# ============================================================

def analyze_1h(df):

    if df.empty or len(df) < 80:
        return None

    df = add_indicators(df)

    df = df.dropna(
        subset=[
            "Tenkan",
            "Kijun",
            "Span_A",
            "Span_B"
        ]
    )

    if len(df) < 2:
        return None

    current = df.iloc[-1]
    previous = df.iloc[-2]

    score = 0

    if current["Close"] < current["Tenkan"]:
        score += 10

    if current["Close"] < current["Kijun"]:
        score += 10

    if current["Tenkan"] < current["Kijun"]:
        score += 10

    if current["Close"] < current["Cloud_Bottom"]:
        score += 15

    if current["Span_A"] < current["Span_B"]:
        score += 10

    if current["Tenkan"] < previous["Tenkan"]:
        score += 5

    if current["Kijun"] < previous["Kijun"]:
        score += 5

    score += volume_score_1h(
        current["Volume_Ratio"]
    )

    return {

        "score": score,

        "close": current["Close"],

        "tenkan": current["Tenkan"],

        "kijun": current["Kijun"],

        "span_a": current["Span_A"],

        "span_b": current["Span_B"],

        "cloud_bottom":
            current["Cloud_Bottom"],

        "volume":
            current["Volume"],

        "volume_avg20":
            current["Volume_Avg20"],

        "volume_ratio":
            current["Volume_Ratio"]
    }


# ============================================================
# ANALYZE STOCK
# ============================================================

def analyze_stock(
    symbol,
    data_5m,
    data_15m,
    data_1h
):

    ticker = symbol + ".NS"

    df5 = get_stock_data(
        data_5m,
        ticker
    )

    df15 = get_stock_data(
        data_15m,
        ticker
    )

    df1h = get_stock_data(
        data_1h,
        ticker
    )

    if df5.empty or df15.empty or df1h.empty:
        return None

    # --------------------------------------------------------
    # REMOVE INCOMPLETE CANDLES
    # --------------------------------------------------------

    df5 = remove_incomplete_last_candle(
        df5,
        5
    )

    df15 = remove_incomplete_last_candle(
        df15,
        15
    )

    df1h = remove_incomplete_last_candle(
        df1h,
        60
    )

    if (
        df5.empty
        or df15.empty
        or df1h.empty
    ):
        return None

    r5 = analyze_5m(df5)
    r15 = analyze_15m(df15)
    r1h = analyze_1h(df1h)

    if (
        r5 is None
        or r15 is None
        or r1h is None
    ):
        return None

    # --------------------------------------------------------
    # STRONG BEARISH STRUCTURE
    # --------------------------------------------------------

    strong_bearish = (

        r1h["close"] < r1h["tenkan"]

        and
        r1h["tenkan"] < r1h["kijun"]

        and
        r15["close"] < r15["tenkan"]

        and
        r15["tenkan"] < r15["kijun"]

        and
        r5["close"] < r5["tenkan"]
    )

    if not strong_bearish:
        return None

    # --------------------------------------------------------
    # FINAL SCORE
    # --------------------------------------------------------

    final_score = (

        r1h["score"] * 0.45

        +
        r15["score"] * 0.35

        +
        r5["score"] * 0.20
    )

    # --------------------------------------------------------
    # VOLUME BONUS
    # --------------------------------------------------------

    avg_volume_ratio = np.nanmean([
        r1h["volume_ratio"],
        r15["volume_ratio"],
        r5["volume_ratio"]
    ])

    volume_bonus = 0

    if avg_volume_ratio >= 2:
        volume_bonus = 5

    elif avg_volume_ratio >= 1.5:
        volume_bonus = 3

    elif avg_volume_ratio >= 1.2:
        volume_bonus = 2

    final_score += volume_bonus

    # --------------------------------------------------------
    # MACD BONUS
    # --------------------------------------------------------

    macd_bonus = 0

    if r5["macd_signal"] == "FRESH SELL":

        age = r5["macd_cross_age"]

        if age == 0:
            macd_bonus = 5

        elif age == 1:
            macd_bonus = 4

        elif age == 2:
            macd_bonus = 3

        elif age == 3:
            macd_bonus = 2

    final_score += macd_bonus

    if final_score < MIN_BEARISH_SCORE:
        return None

    # --------------------------------------------------------
    # RATING
    # --------------------------------------------------------

    if final_score >= 85:
        rating = "VERY STRONG"

    elif final_score >= 70:
        rating = "STRONG"

    elif final_score >= 55:
        rating = "MODERATE"

    else:
        rating = "BEARISH"

    # --------------------------------------------------------
    # CLOUD DISTANCE
    # --------------------------------------------------------

    cloud_bottom = r5["cloud_bottom"]

    if (
        pd.notna(cloud_bottom)
        and cloud_bottom != 0
    ):

        cloud_distance = (
            (
                r5["close"]
                -
                cloud_bottom
            )
            /
            cloud_bottom
        ) * 100

    else:
        cloud_distance = np.nan

    return {

        "Stock": symbol,

        "Final Bearish Score":
            round(final_score, 2),

        "Rating":
            rating,

        "MACD Signal":
            r5["macd_signal"],

        "MACD Cross Time":
            format_time(
                r5["macd_cross_time"]
            ),

        "MACD Cross Age":
            r5["macd_cross_age"],

        "Price":
            round(r5["close"], 2),

        "5M Score":
            r5["score"],

        "15M Score":
            r15["score"],

        "1H Score":
            r1h["score"],

        "5M Volume Ratio":
            round(
                r5["volume_ratio"],
                2
            ),

        "15M Volume Ratio":
            round(
                r15["volume_ratio"],
                2
            ),

        "1H Volume Ratio":
            round(
                r1h["volume_ratio"],
                2
            ),

        "Average Volume Ratio":
            round(
                avg_volume_ratio,
                2
            ),

        "Volume Bonus":
            volume_bonus,

        "MACD Bonus":
            macd_bonus,

        "MACD":
            round(
                r5["macd"],
                4
            ),

        "MACD Signal Value":
            round(
                r5["macd_signal_value"],
                4
            ),

        "MACD Histogram":
            round(
                r5["macd_histogram"],
                4
            ),

        "5M Tenkan":
            round(
                r5["tenkan"],
                2
            ),

        "5M Kijun":
            round(
                r5["kijun"],
                2
            ),

        "5M Cloud Bottom":
            round(
                r5["cloud_bottom"],
                2
            ),

        "5M Cloud Distance %":
            round(
                cloud_distance,
                2
            ),

        "15M Tenkan":
            round(
                r15["tenkan"],
                2
            ),

        "15M Kijun":
            round(
                r15["kijun"],
                2
            ),

        "1H Tenkan":
            round(
                r1h["tenkan"],
                2
            ),

        "1H Kijun":
            round(
                r1h["kijun"],
                2
            ),

        "5M Ichimoku Cross Time":
            format_time(
                r5["cross_time"]
            ),

        "5M Ichimoku Cross Age":
            r5["cross_age"]
    }


# ============================================================
# CREATE EXCEL
# ============================================================

def create_excel(result_df):

    output = BytesIO()

    wb = Workbook()

    ws = wb.active

    ws.title = "Bearish Stocks"

    # --------------------------------------------------------
    # WRITE DATA
    # --------------------------------------------------------

    for col_num, column in enumerate(
        result_df.columns,
        1
    ):

        cell = ws.cell(
            row=1,
            column=col_num,
            value=column
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for row_num, row in enumerate(
        result_df.itertuples(
            index=False
        ),
        2
    ):

        for col_num, value in enumerate(
            row,
            1
        ):

            ws.cell(
                row=row_num,
                column=col_num,
                value=value
            )

    # --------------------------------------------------------
    # FREEZE + FILTER
    # --------------------------------------------------------

    ws.freeze_panes = "A2"

    if ws.max_row >= 1 and ws.max_column >= 1:

        ws.auto_filter.ref = ws.dimensions

    # --------------------------------------------------------
    # STOCK FONT
    # --------------------------------------------------------

    stock_col = None

    for i, cell in enumerate(
        ws[1],
        1
    ):

        if cell.value == "Stock":
            stock_col = i
            break

    if stock_col:

        for row in range(
            2,
            ws.max_row + 1
        ):

            ws.cell(
                row=row,
                column=stock_col
            ).font = Font(
                bold=True
            )

    # --------------------------------------------------------
    # SCORE COLOR SCALE
    # --------------------------------------------------------

    score_col = None

    for i, cell in enumerate(
        ws[1],
        1
    ):

        if cell.value == "Final Bearish Score":
            score_col = i
            break

    if score_col:

        letter = ws.cell(
            row=1,
            column=score_col
        ).column_letter

        ws.conditional_formatting.add(

            f"{letter}2:{letter}{ws.max_row}",

            ColorScaleRule(

                start_type="min",

                start_color="F8696B",

                mid_type="percentile",

                mid_value=50,

                mid_color="FFEB84",

                end_type="max",

                end_color="63BE7B"
            )
        )

    # --------------------------------------------------------
    # MACD SIGNAL FONT
    # --------------------------------------------------------

    macd_col = None

    for i, cell in enumerate(
        ws[1],
        1
    ):

        if cell.value == "MACD Signal":
            macd_col = i
            break

    if macd_col:

        for row in range(
            2,
            ws.max_row + 1
        ):

            cell = ws.cell(
                row=row,
                column=macd_col
            )

            if cell.value == "FRESH SELL":

                cell.font = Font(
                    bold=True
                )

            elif cell.value == "FRESH BUY":

                cell.font = Font(
                    bold=True
                )

    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    for column_cells in ws.columns:

        max_length = 0

        column_letter = (
            column_cells[0]
            .column_letter
        )

        for cell in column_cells:

            try:

                length = len(
                    str(cell.value)
                )

                if length > max_length:
                    max_length = length

            except Exception:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    wb.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# MAIN APP
# ============================================================

st.markdown(
    """
    ### Bearish Scanner

    **5M:** Intraday bearish momentum + fresh MACD SELL  
    **15M:** Intraday Ichimoku confirmation  
    **1H:** Major bearish trend confirmation  
    **Ranking:** 1H Ichimoku → 15M Ichimoku → 5M Ichimoku → Volume → MACD
    """
)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.header("Scanner Settings")

    min_score = st.slider(
        "Minimum Bearish Score",
        min_value=30,
        max_value=100,
        value=MIN_BEARISH_SCORE,
        step=1
    )

    macd_filter = st.selectbox(
        "MACD Filter",
        [
            "ALL",
            "FRESH SELL",
            "FRESH BUY",
            "NONE"
        ]
    )

    st.info(
        "The scanner uses completed "
        "5M, 15M and 1H candles."
    )


# ============================================================
# BUTTONS
# ============================================================

col1, col2 = st.columns(2)

with col1:

    scan_button = st.button(
        "🔄 Scan NIFTY 200",
        use_container_width=True
    )

with col2:

    clear_button = st.button(
        "🗑️ Clear",
        use_container_width=True
    )


if clear_button:

    for key in [
        "sell_results",
        "sell_excel"
    ]:

        if key in st.session_state:
            del st.session_state[key]

    st.rerun()


# ============================================================
# SCAN
# ============================================================

if scan_button:

    try:

        with st.spinner(
            "Loading NIFTY 200..."
        ):

            symbols = get_nifty200()

        tickers = [
            symbol + ".NS"
            for symbol in symbols
        ]

        st.info(
            f"NIFTY 200 stocks loaded: "
            f"{len(symbols)}"
        )

        progress = st.progress(0)

        status = st.empty()

        # ----------------------------------------------------
        # DOWNLOAD 3 TIMEFRAMES
        # ----------------------------------------------------

        status.write(
            "Downloading 5M, 15M and 1H data..."
        )

        with ThreadPoolExecutor(
            max_workers=3
        ) as executor:

            futures = {

                executor.submit(
                    download_data,
                    "5m",
                    PERIOD_5M,
                    tickers
                ): "5m",

                executor.submit(
                    download_data,
                    "15m",
                    PERIOD_15M,
                    tickers
                ): "15m",

                executor.submit(
                    download_data,
                    "60m",
                    PERIOD_1H,
                    tickers
                ): "1h"
            }

            downloaded = {}

            for future in as_completed(
                futures
            ):

                timeframe = futures[
                    future
                ]

                downloaded[
                    timeframe
                ] = future.result()

        data_5m = downloaded["5m"]
        data_15m = downloaded["15m"]
        data_1h = downloaded["1h"]

        # ----------------------------------------------------
        # ANALYZE
        # ----------------------------------------------------

        results = []

        total = len(symbols)

        for i, symbol in enumerate(
            symbols,
            1
        ):

            status.write(
                f"Analyzing {symbol} "
                f"({i}/{total})"
            )

            try:

                result = analyze_stock(
                    symbol,
                    data_5m,
                    data_15m,
                    data_1h
                )

                if result is not None:

                    if (
                        result[
                            "Final Bearish Score"
                        ]
                        >= min_score
                    ):

                        if macd_filter == "ALL":

                            results.append(
                                result
                            )

                        elif (
                            result[
                                "MACD Signal"
                            ]
                            == macd_filter
                        ):

                            results.append(
                                result
                            )

            except Exception:
                pass

            progress.progress(
                i / total
            )

        # ----------------------------------------------------
        # RESULT DATAFRAME
        # ----------------------------------------------------

        result_df = pd.DataFrame(
            results
        )

        if not result_df.empty:

            result_df = (
                result_df
                .sort_values(
                    by=[
                        "Final Bearish Score",
                        "1H Score",
                        "15M Score",
                        "5M Score"
                    ],
                    ascending=False
                )
                .reset_index(
                    drop=True
                )
            )

            excel_bytes = create_excel(
                result_df
            )

            st.session_state[
                "sell_results"
            ] = result_df

            st.session_state[
                "sell_excel"
            ] = excel_bytes

        else:

            st.session_state[
                "sell_results"
            ] = pd.DataFrame()

            st.session_state[
                "sell_excel"
            ] = None

        status.empty()
        progress.empty()

    except Exception as e:

        st.error(
            f"Scanner Error: {e}"
        )


# ============================================================
# DISPLAY RESULTS
# ============================================================

if "sell_results" in st.session_state:

    result_df = st.session_state[
        "sell_results"
    ]

    if result_df is not None and not result_df.empty:

        # ----------------------------------------------------
        # SUMMARY
        # ----------------------------------------------------

        c1, c2, c3, c4 = st.columns(4)

        with c1:

            st.metric(
                "NIFTY 200 Scanned",
                len(get_nifty200())
            )

        with c2:

            st.metric(
                "Bearish Stocks",
                len(result_df)
            )

        with c3:

            st.metric(
                "Fresh MACD SELL",
                int(
                    (
                        result_df[
                            "MACD Signal"
                        ]
                        ==
                        "FRESH SELL"
                    ).sum()
                )
            )

        with c4:

            st.metric(
                "Fresh MACD BUY",
                int(
                    (
                        result_df[
                            "MACD Signal"
                        ]
                        ==
                        "FRESH BUY"
                    ).sum()
                )
            )

        st.success(
            f"{len(result_df)} bearish "
            "stocks found."
        )

        # ----------------------------------------------------
        # TABLE
        # ----------------------------------------------------

        st.dataframe(
            result_df,
            use_container_width=True,
            height=650,
            hide_index=True
        )

        # ----------------------------------------------------
        # EXCEL DOWNLOAD
        # ----------------------------------------------------

        excel_bytes = st.session_state[
            "sell_excel"
        ]

        if excel_bytes:

            st.download_button(
                label="📥 Download Excel",
                data=excel_bytes,
                file_name=(
                    "NIFTY200_Ichimoku_"
                    "Volume_MACD_Bearish.xlsx"
                ),
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
                use_container_width=True
            )

    else:

        st.warning(
            "No bearish stocks matched "
            "the selected conditions."
        )


# ============================================================
# LOGIC EXPANDER
# ============================================================

with st.expander(
    "📋 Scanner Logic"
):

    st.markdown(
        """
        ### 5M Bearish

        - Price < Tenkan
        - Price < Kijun
        - Tenkan < Kijun
        - Price < Cloud Bottom
        - Span A < Span B
        - Tenkan falling
        - Volume confirmation
        - Fresh MACD SELL receives additional score

        ### 15M Bearish

        - Price < Tenkan
        - Price < Kijun
        - Tenkan < Kijun
        - Price < Cloud Bottom
        - Span A < Span B
        - Tenkan falling
        - Kijun falling
        - Volume confirmation

        ### 1H Bearish

        - Price < Tenkan
        - Price < Kijun
        - Tenkan < Kijun
        - Price < Cloud Bottom
        - Span A < Span B
        - Tenkan falling
        - Kijun falling
        - Strong volume confirmation

        ### Strong Bearish Structure

        The stock must satisfy:

        **1H**
        - Close < Tenkan
        - Tenkan < Kijun

        **15M**
        - Close < Tenkan
        - Tenkan < Kijun

        **5M**
        - Close < Tenkan

        ### Final Score

        - 1H = 45%
        - 15M = 35%
        - 5M = 20%
        - Volume bonus
        - Fresh MACD SELL bonus

        ### Rating

        - 85+ = VERY STRONG
        - 70+ = STRONG
        - 55+ = MODERATE
        - Below 55 = BEARISH
        """
    )
```

### `requirements.txt`

```text
streamlit
yfinance
pandas
numpy
requests
openpyxl
```

### Excel output

The app creates:

**`NIFTY200_Ichimoku_Volume_MACD_Bearish.xlsx`**

with sheet:

**`Bearish Stocks`**

The **highest bearish score is placed first**, Excel filters are enabled, the first row is frozen, and the workbook is timezone-safe.

This is the bearish mirror of the previous scanner; importantly, it uses **FRESH MACD SELL** rather than treating every MACD-below-signal condition as a fresh sell.
